"""
Reportes: Excel (para trabajar/compartir) y HTML interactivo autocontenido
(sin CDNs: abre sin internet y se puede mandar por correo).
"""
import json
import math
import os

import numpy as np
import pandas as pd

from . import metricas as M
from . import periodos as P


def _limpio(x):
    if isinstance(x, dict):
        return {str(k): _limpio(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_limpio(v) for v in x]
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating, float)):
        return None if (math.isnan(x) or math.isinf(x)) else float(x)
    if isinstance(x, (np.bool_,)):
        return bool(x)
    if x is pd.NaT:
        return None
    return x


def _registros(df):
    if df is None or len(df) == 0:
        return []
    return _limpio(df.replace({np.nan: None}).to_dict("records"))


def _orden_bins(bdef, base_counts):
    if bdef and bdef.get("clase") == "numerica":
        k = len(bdef.get("cortes") or []) + 1
        return ["b%02d" % i for i in range(1, k + 1)] + ["__INF_NAN__", "__NULO__"]
    cats = sorted([b for b in base_counts if b not in ("__OTROS__", "__NULO__")], key=lambda b: -base_counts[b])
    return cats + ["__OTROS__", "__NULO__"]


def _esperado_payload(r):
    defs = r.get("psi_esp_def") or {}
    besp = r.get("besp")
    out = {}
    for v, d in defs.items():
        counts = {}
        if besp is not None and not besp.empty:
            for p, gp in besp[besp["variable"] == v].groupby("periodo"):
                counts[str(int(p))] = {str(b): int(n) for b, n in zip(gp["bin"], gp["n"])}
        if d["clase"] == "numerica":
            orden = ["b%02d" % (i + 1) for i in range(len(d["cortes"]) + 1)] + ["__INF_NAN__", "__NULO__"]
            et = M.etiquetas_bins(d["cortes"])
        else:
            orden = list(d["categorias"]) + ["__OTROS__", "__NULO__"]
            et = {"__OTROS__": "(otros)", "__NULO__": "NULO"}
        out[v] = {"orden": orden, "etiquetas": et, "dist": d["esperado"], "counts": counts}
    return out


def _lectura_payload(lec):
    if not lec:
        return None
    info = dict(lec["info"])
    hojas = info.pop("hojas", [])
    return {"info": [{"k": k, "v": ", ".join(map(str, v)) if isinstance(v, list) else str(v)} for k, v in info.items()],
            "inferencia": _registros(lec["inferencia"]), "hojas": hojas}


def construir_payload(r):
    g = r["campo"]["granularidad"]
    periodos = [int(p) for p in r["periodos"]]
    ini, fin = r["rango"]
    rango = P.rango_periodos(ini, fin, g)
    etiquetas = {str(p): P.etiqueta(p, g) for p in rango}

    st = r["stats"]
    fo = r.get("forma")
    if st is not None and not st.empty and fo is not None and not fo.empty:
        st = st.merge(fo[["periodo", "variable", "pct_fuera_15", "pct_fuera_3", "pct_centinela", "asimetria",
                          "curtosis_exceso"]], on=["periodo", "variable"], how="left")
    stats = {}
    if st is not None and not st.empty:
        keep = [c for c in st.columns if c not in ("variable", "tipo")]
        for v, gdf in st.groupby("variable"):
            stats[v] = _registros(gdf[keep].sort_values("periodo"))
    psi = {}
    if r["psi"] is not None and not r["psi"].empty:
        for v, gdf in r["psi"].groupby("variable"):
            psi[v] = _registros(gdf.drop(columns=["variable"]))
    bins = {}
    if r["bins"] is not None and not r["bins"].empty:
        base = r["periodo_base"]
        for v, gdf in r["bins"].groupby("variable"):
            counts = {}
            for p, gp in gdf.groupby("periodo"):
                counts[str(int(p))] = {str(b): int(n) for b, n in zip(gp["bin"], gp["n"])}
            bdef = r["bins_def"].get(v, {})
            orden = _orden_bins(bdef, counts.get(str(base), {}))
            et = M.etiquetas_bins(bdef.get("cortes") or []) if bdef.get("clase") == "numerica" else {
                "__OTROS__": "(otros)", "__NULO__": "NULO"}
            bins[v] = {"orden": orden, "etiquetas": et, "counts": counts}

    cols = r["columnas"].copy()
    rol = {}
    for c in r["campo"]["columnas"]:
        rol[c] = "campo fecha"
    for c in r.get("cols_particion") or []:
        rol.setdefault(c, "partición")
    for c in r.get("llave") or []:
        rol[c.lower()] = "llave"
    perfiladas = {v for v, _ in r["variables"]}
    cols["rol"] = [rol.get(c, "perfilada" if c in perfiladas else "excluida") for c in cols["columna"]]

    parts = r["particiones"]
    parts_rec = _registros(parts.head(3000)) if parts is not None and not parts.empty else []

    modo_txt = {"rango": "rango %s a %s", "puntual": "puntual (%s vs %s)", "todo": "todo el histórico (%s a %s)"}[r["modo"]]
    modo_txt = modo_txt % ((P.etiqueta(ini, g), P.etiqueta(fin, g)) if r["modo"] != "puntual" else (P.etiqueta(fin, g), P.etiqueta(ini, g)))
    tareas = r["tareas"]
    params = [
        ("tabla", r["tabla"]), ("campo fecha", ",".join(r["campo"]["columnas"])), ("tipo campo", r["campo"]["tipo"]),
        ("granularidad", g), ("modo", modo_txt), ("periodo base PSI", P.etiqueta(r["periodo_base"], g)),
        ("llave", ",".join(r["llave"] or []) or "-"), ("filtro extra", r.get("filtro_extra") or "-"),
        ("zona temporal", r["zona_temporal"]), ("n_tiles (percentiles)", r["n_tiles"]), ("n_bins PSI", r["n_bins"]),
    ] + [("umbral " + k, v) for k, v in r["umbrales"].items()]

    return _limpio({
        "meta": {
            "tabla": r["tabla"], "run_key": r["run_key"], "campo": ",".join(r["campo"]["columnas"]),
            "tipo_campo": r["campo"]["tipo"], "granularidad": g, "modo_txt": modo_txt,
            "periodo_base": int(r["periodo_base"]), "fecha": r["fecha_ejecucion"], "duracion": r.get("duracion_seg"),
            "n_tareas": int(len(tareas)), "n_fallidas": int(len(r["fallidas"])), "umbrales": r["umbrales"],
            "filtro_extra": r.get("filtro_extra"), "n_cols_part": len(r.get("cols_particion") or []),
            "foto": bool((r.get("lectura") or {}).get("info", {}).get("foto_unica")),
            "origen": "archivo" if r.get("lectura") else "tabla",
        },
        "periodos": periodos, "rango_periodos": rango, "etiquetas": etiquetas, "huecos": r["huecos"],
        "volumen": _registros(r["volumen"]), "uniq": _registros(r["uniq"]),
        "vars": _registros(r["resumen_variables"]), "stats": stats, "psi": psi, "bins": bins,
        "alertas": _registros(r["alertas"]), "columnas": _registros(cols), "particiones": parts_rec,
        "tareas": _registros(tareas.sort_values(["periodo", "fase"]) if not tareas.empty else tareas),
        "parametros": [{"k": k, "v": str(v)} for k, v in params],
        "reglas": _registros(r.get("reglas")), "cobertura": _registros(r.get("cobertura")),
        "texto": _registros(r.get("texto")), "fechas": _registros(r.get("fechas")),
        "regimen": _registros(r.get("regimen")), "otros": _registros(r.get("otros")),
        "centinelas": _registros(r.get("centinelas")), "identicas": _registros(r.get("identicas")),
        "lectura": _lectura_payload(r.get("lectura")),
        "esperado": _esperado_payload(r),
        "vista_sql": r.get("vista_limpia_sql", ""), "vista_archivo": os.path.basename(r.get("sql", "") or ""),
    })


def generar_html(r, ruta):
    plantilla = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "reporte.html")
    with open(plantilla, encoding="utf-8") as f:
        html = f.read()
    datos = json.dumps(construir_payload(r), ensure_ascii=False, allow_nan=False).replace("</", "<\\/")
    html = html.replace("/*__DATOS__*/null", datos)
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(html)
    return ruta


def generar_excel(r, ruta):
    g = r["campo"]["granularidad"]
    resumen = pd.DataFrame([
        ("Tabla", r["tabla"]), ("Campo fecha", ",".join(r["campo"]["columnas"])), ("Tipo campo", r["campo"]["tipo"]),
        ("Granularidad", g), ("Rango", "%s a %s" % (P.etiqueta(r["rango"][0], g), P.etiqueta(r["rango"][1], g))),
        ("Periodo base PSI", P.etiqueta(r["periodo_base"], g)), ("Periodos con datos", len(r["periodos"])),
        ("Huecos", ", ".join(P.etiqueta(p, g) for p in r["huecos"]) or "-"),
        ("Variables", len(r["variables"])), ("Tareas fallidas", len(r["fallidas"])),
        ("Alertas ALTA", int((r["alertas"]["nivel"] == "ALTA").sum()) if not r["alertas"].empty else 0),
        ("Fecha ejecución", r["fecha_ejecucion"]), ("Run key", r["run_key"]),
    ], columns=["concepto", "valor"])
    psi_mat = pd.DataFrame()
    if r["psi"] is not None and not r["psi"].empty:
        psi_mat = r["psi"].pivot_table(index="variable", columns="periodo", values="psi_anterior")
    nul_mat = pd.DataFrame()
    if r["stats"] is not None and not r["stats"].empty:
        nul_mat = r["stats"].pivot_table(index="variable", columns="periodo", values="pct_nulos")

    hojas = [
        ("resumen", resumen, False), ("alertas", r["alertas"], False), ("ficha_variables", r["resumen_variables"], False),
        ("volumen", r["volumen"], False), ("psi_matriz", psi_mat, True), ("nulos_matriz", nul_mat, True),
        ("estadisticos", r["stats"], False), ("psi_detalle", r["psi"], False), ("buckets", r["bins"], False),
        ("unicidad", r["uniq"], False), ("columnas", r["columnas"], False), 
        ("outliers_forma", r.get("forma"), False), ("reglas", r.get("reglas"), False),
        ("centinelas", r.get("centinelas"), False), ("texto_tipos", r.get("texto"), False),
        ("fechas", r.get("fechas"), False), ("cobertura", r.get("cobertura"), False),
        ("regimen", r.get("regimen"), False), ("categorias_otros", r.get("otros"), False),
        ("columnas_identicas", r.get("identicas"), False),
        ("lectura_tipos", (r.get("lectura") or {}).get("inferencia"), False),
        ("tareas", r["tareas"], False),
    ]
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for nombre, df, idx in hojas:
            if df is None or len(df) == 0:
                pd.DataFrame({"info": ["sin datos"]}).to_excel(xw, sheet_name=nombre, index=False)
                continue
            df.to_excel(xw, sheet_name=nombre, index=idx)
        _formato(xw, r)
    return ruta


def _formato(xw, r):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    u = r["umbrales"]
    for ws in xw.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "B2"
        for i, col in enumerate(ws.columns, start=1):
            largo = max((len(str(c.value)) for c in list(col)[:200] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(9, largo + 2))
    if "psi_matriz" in xw.book.sheetnames and xw.book["psi_matriz"].max_column >= 2 and xw.book["psi_matriz"].max_row >= 2:
        ws = xw.book["psi_matriz"]
        rng = "B2:%s%d" % (get_column_letter(ws.max_column), ws.max_row)
        for op, val, color in (("greaterThanOrEqual", u["psi_alto"], "F4B6B6"),
                               ("between", (u["psi_moderado"], u["psi_alto"]), "F9D2B8"),
                               ("lessThan", u["psi_moderado"], "CDEFCB")):
            formula = [str(v) for v in val] if isinstance(val, tuple) else [str(val)]
            ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=formula,
                                                          fill=PatternFill("solid", fgColor=color)))
    if "nulos_matriz" in xw.book.sheetnames and xw.book["nulos_matriz"].max_column >= 2 and xw.book["nulos_matriz"].max_row >= 2:
        ws = xw.book["nulos_matriz"]
        rng = "B2:%s%d" % (get_column_letter(ws.max_column), ws.max_row)
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                          end_type="num", end_value=1, end_color="2A78D6"))
        for row in ws.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.number_format = "0.0%"


def generar_reportes(r, dir_run):
    base = "perfil_%s_%s" % (r["tabla"].replace(".", "__"), r["run_key"])
    rutas = {}
    if r.get("vista_limpia_sql"):
        rutas["sql"] = os.path.join(dir_run, "vista_limpia_%s.sql" % base.replace("perfil_", ""))
        with open(rutas["sql"], "w", encoding="utf-8") as f:
            f.write(r["vista_limpia_sql"])
        r["sql"] = rutas["sql"]
    rutas["html"] = generar_html(r, os.path.join(dir_run, base + ".html"))
    try:
        rutas["excel"] = generar_excel(r, os.path.join(dir_run, base + ".xlsx"))
    except Exception as exc:  # p.ej. sin openpyxl
        if "openpyxl" in str(exc):
            print("Excel no generado: falta el paquete openpyxl (instálelo con: pip install openpyxl). "
                  "El reporte HTML sí se generó.")
        else:
            print("No se pudo generar Excel: %s" % exc)
    return rutas
